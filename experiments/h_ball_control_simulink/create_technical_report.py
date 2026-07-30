from pathlib import Path
import math
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "钢球水管平衡系统算法与Simulink仿真说明.docx"
FLOW_IMAGE = ROOT / "control_algorithm_flow.png"

BLUE = "2E74B5"
DARK_BLUE = "0B2545"
DEEP_BLUE = "1F4D78"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
CAUTION = "FFF4CE"
RISK = "FDECEC"
GREEN = "EAF5EC"
WHITE = "FFFFFF"
BLACK = "111827"


def set_run_font(run, *, latin="Calibri", east_asia="Microsoft YaHei",
                 size=None, color=None, bold=None, italic=None):
    run.font.name = latin
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    fonts = run._element.rPr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start),
                          ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def configure_table(table, widths, header=True, font_size=9):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for r_idx, row in enumerate(table.rows):
        if header and r_idx == 0:
            tr_pr = row._tr.get_or_add_trPr()
            tbl_header = tr_pr.find(qn("w:tblHeader"))
            if tbl_header is None:
                tbl_header = OxmlElement("w:tblHeader")
                tr_pr.append(tbl_header)
            tbl_header.set(qn("w:val"), "true")
        for c_idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[c_idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if header and r_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=font_size,
                        bold=(header and r_idx == 0),
                        color=BLACK,
                    )


def add_table(doc, headers, rows, widths, alignments=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = str(header)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            cells[idx].text = str(value)
    configure_table(table, widths, header=True, font_size=font_size)
    if alignments:
        for row in table.rows:
            for idx, alignment in enumerate(alignments):
                row.cells[idx].paragraphs[0].alignment = alignment
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_left_border(paragraph, color=BLUE, size=18, space=8):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    borders.append(left)


def add_callout(doc, label, text, fill=CALLOUT, color=DEEP_BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.06)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.10
    add_paragraph_shading(p, fill)
    add_left_border(p, color=color)
    r = p.add_run(f"{label}  ")
    set_run_font(r, size=10.5, color=color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=BLACK)
    return p


def add_equation(doc, text):
    p = doc.add_paragraph()
    p.style = doc.styles["Equation"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(
        r, latin="Cambria Math", east_asia="Microsoft YaHei",
        size=10.5, color=DARK_BLUE
    )
    return p


def add_figure(doc, image_path, caption, width=6.3):
    image_path = Path(image_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(image_path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption.split("  ", 1)[0])
    c = doc.add_paragraph()
    c.style = doc.styles["Caption"]
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(8)
    r = c.add_run(caption)
    set_run_font(r, size=9, color=MUTED)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def create_numbering(doc):
    numbering = doc.part.numbering_part.element

    def add_definition(abstract_id, num_id, num_fmt, text, marker_font=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_fmt)
        lvl.append(fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        if marker_font:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), marker_font)
            fonts.set(qn("w:hAnsi"), marker_font)
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)

    add_definition(90, 90, "bullet", "•", "Symbol")
    add_definition(91, 91, "decimal", "%1.")
    return 90, 91


def add_list_item(doc, text, marker):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.paragraph_format.left_indent = Inches(0.40)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f"{marker}  ")
    set_run_font(r, size=11, color=BLACK)
    r = p.add_run(text)
    set_run_font(r, size=11, color=BLACK)
    return p


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DEEP_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0

    if "Equation" not in styles:
        eq = styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        eq = styles["Equation"]
    eq.font.name = "Cambria Math"
    eq.font.size = Pt(10.5)
    eq._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    eq.paragraph_format.space_before = Pt(3)
    eq.paragraph_format.space_after = Pt(6)
    eq.paragraph_format.keep_together = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run("钢球—水管平衡系统  |  技术说明报告")
    set_run_font(r, size=9, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("第 ")
    set_run_font(r, size=9, color=MUTED)
    add_page_field(fp)
    r = fp.add_run(" 页")
    set_run_font(r, size=9, color=MUTED)

    first_footer = section.first_page_footer
    ffp = first_footer.paragraphs[0]
    ffp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ffp.add_run("设计仿真版 · 参数待实测标定")
    set_run_font(r, size=9, color=MUTED)


def draw_arrow(draw, start, end, fill, width=5):
    draw.line([start, end], fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 16
    p1 = (
        end[0] - size * math.cos(angle - math.pi / 6),
        end[1] - size * math.sin(angle - math.pi / 6),
    )
    p2 = (
        end[0] - size * math.cos(angle + math.pi / 6),
        end[1] - size * math.sin(angle + math.pi / 6),
    )
    draw.polygon([end, p1, p2], fill=fill)


def create_flow_image():
    width, height = 1800, 980
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "msyh.ttc"
    bold_path = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "msyhbd.ttc"
    font = ImageFont.truetype(str(font_path), 30)
    small = ImageFont.truetype(str(font_path), 23)
    title = ImageFont.truetype(str(bold_path), 35)

    def box(rect, heading, lines, fill, outline=DEEP_BLUE):
        x1, y1, x2, y2 = rect
        draw.rounded_rectangle(
            rect, radius=22, fill=f"#{fill}", outline=f"#{outline}", width=4
        )
        bbox = draw.textbbox((0, 0), heading, font=title)
        tx = (x1 + x2 - (bbox[2] - bbox[0])) / 2
        draw.text((tx, y1 + 20), heading, fill=f"#{DARK_BLUE}", font=title)
        y = y1 + 78
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=small)
            tx = (x1 + x2 - (bbox[2] - bbox[0])) / 2
            draw.text((tx, y), line, fill=f"#{BLACK}", font=small)
            y += 38

    box((55, 100, 375, 285), "树莓派", ["钢球位置", "100 Hz + 时间戳"], LIGHT_BLUE)
    box((55, 390, 375, 575), "MSPM0", ["车身IMU源样本", "约30 Hz + 时间戳"], LIGHT_BLUE)
    box((55, 680, 375, 865), "RS00反馈", ["角度 / 速度", "力矩 / 温度"], LIGHT_BLUE)

    box(
        (545, 125, 1045, 835),
        "PSoC Edge E84 · M33",
        [
            "时间同步与姿态/加速度处理",
            "",
            "延迟卡尔曼：x、v、扰动d",
            "",
            "LQI + IMU前馈 + 抗积分饱和",
            "",
            "±6°限角 / 0.35 rad/s限速",
            "",
            "两连杆逆解与失联保护",
            "",
            "200 Hz控制 / 500 Hz安全执行",
        ],
        "F4F6F9",
    )

    box((1215, 145, 1715, 350), "RS00执行器", ["CAN运动模式", "主动杆35.0 mm"], GREEN)
    box((1215, 455, 1715, 660), "两连杆与水管", ["蓝杆55.5 mm", "水管摇杆300.1 mm"], GREEN)
    box((1215, 765, 1715, 925), "钢球对象", ["滚动 / 滑动", "摩擦与车辆扰动"], GREEN)

    for y in (192, 482, 772):
        draw_arrow(draw, (375, y), (545, y), f"#{BLUE}")
    draw_arrow(draw, (1045, 245), (1215, 245), f"#{BLUE}")
    draw_arrow(draw, (1465, 350), (1465, 455), f"#{BLUE}")
    draw_arrow(draw, (1465, 660), (1465, 765), f"#{BLUE}")
    draw_arrow(draw, (1215, 845), (1075, 845), f"#{BLUE}")
    draw.line([(1075, 845), (1075, 900), (435, 900), (435, 192)], fill=f"#{BLUE}", width=5)
    draw_arrow(draw, (435, 192), (545, 192), f"#{BLUE}")

    draw.text((1110, 790), "位置反馈", fill=f"#{MUTED}", font=font)
    draw.text((60, 35), "图：实物控制链与算法闭环", fill=f"#{DARK_BLUE}", font=title)
    image.save(FLOW_IMAGE, dpi=(180, 180))


def add_body_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True, color=DARK_BLUE)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, color=BLACK)
    else:
        r = p.add_run(text)
        set_run_font(r, color=BLACK)
    return p


def build_report():
    create_flow_image()
    doc = Document()
    style_document(doc)

    # Cover: editorial report pattern, using the standard_business_brief preset.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("技术说明报告")
    set_run_font(r, size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("钢球—水管两连杆平衡系统")
    set_run_font(r, size=28, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("机械原理、控制算法与Simulink非线性仿真")
    set_run_font(r, size=15, color=DEEP_BLUE)

    add_callout(
        doc,
        "适用方案",
        "RS00关节电机直接驱动35.0 mm主动杆，经55.5 mm连杆带动300.1 mm水管摇杆；"
        "水管有效滚动段250 mm。PSoC Edge E84直接控制RS00，MSPM0发布车身IMU，树莓派提供视觉位置。",
        fill=LIGHT_BLUE,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(38)
    p.paragraph_format.space_after = Pt(5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("版本：V1.3（实测连杆、实机视觉与约30 Hz IMU基线）")
    set_run_font(r, size=11, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("日期：2026年7月31日")
    set_run_font(r, size=11, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("状态：机构尺寸已形成建议值，摩擦/延迟/1 cm性能仍需实物标定")
    set_run_font(r, size=10.5, color=MUTED)

    doc.add_page_break()

    doc.add_heading("目录", level=1)
    toc_items = [
        "系统问题与结论概览",
        "系统构成：哪些属于机械，哪些属于算法",
        "两连杆机械运动学与支撑高度",
        "钢球在水管中的动力学",
        "传感器融合与状态估计",
        "LQI控制、车辆前馈与安全监督",
        "Simulink非线性仿真模型",
        "仿真结果与压力测试解释",
        "英飞凌部署方案",
        "实物参数辨识与验证",
        "结论与当前设计基线",
        "附录：模型文件与关键参数",
    ]
    for index, item in enumerate(toc_items, start=1):
        add_list_item(doc, item, f"{index}.")
    add_callout(
        doc,
        "阅读提示",
        "本文中的“仿真通过”只说明给定模型和参数下没有掉球或越界，不自动等同于实物满足"
        "误差小于1 cm。实物结论必须建立在参数辨识、硬件在环和重复道路测试之上。",
        fill=CAUTION,
        color="7A5A00",
    )
    doc.add_page_break()

    doc.add_heading("1 系统问题与结论概览", level=1)
    add_body_paragraph(
        doc,
        "本系统的目标，是在车辆运动、视觉延迟、低摩擦和机构非线性的共同作用下，"
        "通过改变水管倾角，使钢球保持在指定位置，并在题目规定的评价区间内把位置误差控制在10 mm以内。"
    )
    add_callout(
        doc,
        "核心结论",
        "这不是单纯的机械问题，也不是只换一个PID就能解决的问题。机构几何决定可用倾角和传动比；"
        "钢球动力学决定控制对象；卡尔曼负责估计位置、速度和扰动；LQI负责闭环调节；"
        "IMU前馈负责抵消车辆加减速；安全监督器负责在管端和通信异常时覆盖普通控制。",
        fill=LIGHT_BLUE,
    )
    add_table(
        doc,
        ["项目", "当前设计基线", "性质"],
        [
            ("水管有效滚动长度 / CB", "250 mm / 300.1 mm", "用户给定；两者不是同一个量"),
            ("主动杆 / 蓝色连杆", "35.0 mm / 55.5 mm", "用户给定"),
            ("RS00轴心O高度", "38 mm", "用户给定"),
            ("水管合页轴C高度 / O-C水平距", "93 mm / 285 mm", "用户给定"),
            ("对称软件硬限角", "暂定±6°", "新机构几何安全值，待无球测试确认"),
            ("IMU唯一源样本率", "约29.1 Hz", "当前WIT 9600 bit/s理论上限"),
            ("外环控制频率", "200 Hz", "设计值"),
            ("视觉频率", "100 Hz", "用户最新确认"),
            ("位置精度", "评价区间内峰值≤10 mm", "题目要求，需按原文确定时间窗"),
        ],
        [2250, 3150, 3960],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_heading("2 系统构成：哪些属于机械，哪些属于算法", level=1)
    add_figure(
        doc,
        FLOW_IMAGE,
        "图1  实物控制链：PSoC E84完成融合、控制与RS00安全执行；MSPM0只承担IMU采集。",
        width=6.3,
    )
    add_body_paragraph(
        doc,
        "机械原理部分包括两连杆几何约束、角度传动比、虚功力矩换算、等效转动惯量，以及钢球的"
        "平动、转动、滚动和滑动。它们决定“给RS00一个角度后，水管和钢球实际上怎样运动”。"
    )
    add_body_paragraph(
        doc,
        "控制算法部分包括卡尔曼状态估计、LQI/LQR反馈、IMU前馈、延迟测量处理、积分抗饱和和"
        "端部安全监督。它们决定“根据观测到的状态，下一时刻应该把水管倾到多少角度”。"
    )
    add_table(
        doc,
        ["方法", "类别", "在本项目中的作用", "是否直接控制电机"],
        [
            ("两连杆正/逆运动学", "机械运动学", "水管角与RS00角之间转换", "间接"),
            ("钢球滚动/滑动模型", "机械动力学", "预测倾角对钢球加速度的作用", "否"),
            ("卡尔曼滤波", "状态估计", "融合延迟视觉、角度与IMU，估计x、v、扰动", "否"),
            ("LQI/LQR", "反馈控制", "根据位置、速度和积分误差计算目标倾角", "是"),
            ("ADRC/ESO", "鲁棒控制/观测", "可估计未建模扰动，但不宜与现有扰动状态重复堆叠", "可选"),
            ("车辆前馈", "补偿控制", "在钢球明显偏移前抵消车辆加速和俯仰", "是"),
            ("MPC", "预测优化", "更适合地图速度和加减速规划，不作为首版内环", "间接"),
        ],
        [1500, 1320, 4400, 2140],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
         WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
        font_size=8.7,
    )

    doc.add_heading("3 两连杆机械运动学与支撑高度", level=1)
    doc.add_heading("3.1 几何定义", level=2)
    add_body_paragraph(
        doc,
        "以右侧水管固定端合页轴C为原点。O为左侧RS00轴心，A为主动杆末端，B为蓝色连杆与"
        "水管的连接点。模型采用O=(−285,−55) mm；水管水平时B=(−300.1,0) mm，即C到B向左。"
    )
    add_equation(doc, "A = O + r[cos(q + q₀), sin(q + q₀)]ᵀ")
    add_equation(doc, "B = C + L[cos(θ+π), sin(θ+π)]ᵀ，且  ‖B − A‖ = l_AB")
    add_body_paragraph(
        doc,
        "其中r=35.0 mm、l=55.5 mm、L=300.1 mm，q为RS00角度，θ为相对车身的水管角，"
        "q₀为装配零点偏置。250 mm只表示钢球有效滚动段，不能代替CB进入四杆方程。"
        "每次控制先求期望水管角θ，再通过逆运动学求RS00目标角q。"
    )
    add_figure(
        doc,
        ROOT / "two_link_mechanism_analysis.png",
        "图2  两连杆机构仿真分析：−6°、0°、+6°构型，逆解、速度/力矩倍率和等效惯量随姿态变化。",
        width=6.3,
    )

    doc.add_heading("3.2 C轴高度93 mm与工作区", level=2)
    add_equation(doc, "H_C = 93 mm，H_O = 38 mm，ΔH = 55 mm，ΔX_OC = 285 mm")
    add_body_paragraph(
        doc,
        "真正决定连杆几何的是C轴和O轴的相对位置，不是支撑块外形高度。水管水平时，B相对O为"
        "(−15.1,55) mm，因此OB≈57.0 mm，满足两圆相交条件20.5 mm<OB<90.5 mm。"
    )
    add_callout(
        doc,
        "加工基准",
        "机械图应标注“右侧水管合页C轴中心距底板93 mm、C轴相对O轴水平285 mm”。若水管"
        "外径为50 mm且C轴与管中心重合，管底约68 mm，满足支撑台不低于50 mm的约束。",
        fill=GREEN,
        color="2F6B3A",
    )
    add_body_paragraph(
        doc,
        "在当前|dθ/dq|≥0.03保护下，水平装配分支连续区约为−6.47°至+6.79°。因此取对称±6°"
        "作为暂定软件硬指令限位，并把正常反馈限制在±4°、端部恢复限制在±5.5°。"
    )
    add_equation(doc, "20.5 mm ≤ ‖B(θ) − O‖ ≤ 90.5 mm")
    add_body_paragraph(
        doc,
        "在±6°范围内，RS00相对水平零位约需运动−57.00°至+70.68°；最小|dθ/dq|=0.0467，"
        "最不利速度/力矩换算倍率约21.43。水管指令斜率因此从旧值1.2 rad/s降为0.35 rad/s。"
    )

    doc.add_heading("4 钢球在水管中的动力学", level=1)
    doc.add_heading("4.1 纯滚动的基本规律", level=2)
    add_body_paragraph(
        doc,
        "钢球不是质点。倾斜水管时，重力既产生平动，也通过接触摩擦产生转动。实心球转动惯量为"
        "I=2mr²/5，在不打滑的理想情况下："
    )
    add_equation(doc, "ẍ = (5/7)g sinθ")
    add_body_paragraph(
        doc,
        "因此1°水管角就能产生约0.122 m/s²的钢球加速度。持续0.4 s的角度偏差已经可能造成接近"
        "10 mm的位置变化，这也是机构角度标定和低延迟控制非常重要的原因。"
    )

    doc.add_heading("4.2 滚动、微滑与纯滑动", level=2)
    add_body_paragraph(
        doc,
        "实物中钢球会在纯滚动、微滑和明显滑动之间切换。仿真使用相对滑动速度vₛ=v−rω，并以"
        "Stribeck形式描述静摩擦到动摩擦的连续过渡："
    )
    add_equation(doc, "μ(vₛ) = μ_k + (μ_s − μ_k)exp[−(|vₛ|/v_s)²]")
    add_body_paragraph(
        doc,
        "模型还加入滚动阻力矩、黏性阻力、速度平方阻力和接触法向力变化。当前摩擦参数只是"
        "保守假设，必须通过斜管启动角和不同倾角下的位置—时间曲线进行辨识。"
    )

    doc.add_heading("4.3 车辆运动产生的非惯性力", level=2)
    add_equation(
        doc,
        "F_axis = m[(g+a_z)sin(θ+φ) − a_x cos(θ+φ)]",
    )
    add_body_paragraph(
        doc,
        "其中φ为车体俯仰，aₓ为沿水管方向的车辆加速度，a_z为竖直振动。车辆向前加速会在车体"
        "坐标系中把钢球推向相反方向；因此仅靠相机看到球已经移动后再纠正，会产生明显滞后，必须使用IMU前馈。"
    )
    add_callout(
        doc,
        "车身IMU能做什么",
        "IMU装在车身上正适合测量底座俯仰和车辆加速度，用于前馈补偿。它不能测量水管相对车身"
        "的倾角，也看不到连杆回差和支架变形；相对水管角必须由RS00编码器经过四杆正解得到，"
        "必要时在C轴增加直接角度编码器。",
        fill=LIGHT_BLUE,
    )

    doc.add_heading("5 传感器融合与状态估计", level=1)
    doc.add_heading("5.1 为什么不能只用相机位置做PID", level=2)
    add_body_paragraph(
        doc,
        "相机为100 Hz，但仍包含曝光、树莓派处理和串口传输延迟。直接对相机位置差分会放大像素噪声，"
        "直接把旧画面当成当前位置又会造成相位滞后。低摩擦钢球惯性明显，控制器真正需要的是当前位置、"
        "速度和外界扰动，而不是一帧孤立的位置。"
    )
    add_table(
        doc,
        ["信息", "来源", "模型频率/延迟", "用途"],
        [
            ("钢球位置", "树莓派视觉", "100 Hz；默认35 ms", "校正位置状态"),
            ("车体俯仰/轴向加速度", "MSPM0 IMU", "约30 Hz唯一源样本；默认35 ms", "保持输入、车辆前馈"),
            ("RS00角度/速度", "RS00 CAN", "模型500 Hz、约3.5 ms；实机待测", "计算真实水管角及安全"),
            ("目标位置", "题目/地图规划", "200 Hz控制读取", "通常为水管中心"),
        ],
        [1800, 2200, 2500, 2860],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
         WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_callout(
        doc,
        "115200波特率升级评估",
        "三类11字节8N1帧每个完整组至少330 bit。9600 bit/s理论上限约29.1组/s；"
        "115200 bit/s理论上限约349.1组/s，可支持100 Hz或200 Hz完整组，但仍不能支持500 Hz。"
        "同一可行综合压力场景中，30 Hz基线RMS/峰值为17.78/45.73 mm；假设200 Hz唯一源样本"
        "和8 ms延迟后降为10.62/25.00 mm。两端波特率和WIT输出率必须一起修改。",
        fill=GREEN,
        color="2F6B3A",
    )

    doc.add_heading("5.2 增广卡尔曼状态", level=2)
    add_equation(doc, "z = [x, v, d]ᵀ")
    add_body_paragraph(
        doc,
        "x为钢球位置，v为速度，d为摩擦变化、模型误差等合并扰动加速度。控制器每5 ms执行一次预测；"
        "新相机帧和新IMU源样本只按各自源序号使用一次。当前200 Hz CAN姿态镜像中的重复值只做保持，"
        "不能刷新IMU新鲜度。"
    )
    add_equation(
        doc,
        "θ_eff = θ_pipe + φ − a_x/g，    ẍ ≈ (5/7)g·θ_eff + d",
    )
    add_body_paragraph(
        doc,
        "现有Simulink使用延迟观测矩阵近似处理旧位置。实物部署建议保存最近100–150 ms状态环形缓冲，"
        "在拍摄时刻更新后重新预测到当前时刻，可以更准确地处理变化的视觉延迟。"
    )

    doc.add_heading("6 LQI控制、车辆前馈与安全监督", level=1)
    doc.add_heading("6.1 控制律", level=2)
    add_equation(
        doc,
        "θ_cmd = a_x/g − φ − d̂/b − K_p(x̂−r) − K_vv̂ − K_iξ",
    )
    add_equation(doc, "ξ(k+1) = ξ(k) + T_s[x̂(k) − r(k)]，    b = 5g/7")
    add_body_paragraph(
        doc,
        "第一、第二项是车辆加速度与俯仰前馈；第三项补偿估计扰动；后三项是LQI反馈。输出再经过"
        "±6°硬限角、0.35 rad/s斜率限制、积分抗饱和、两连杆逆解和RS00安全限制。"
    )
    add_callout(
        doc,
        "算法选择",
        "卡尔曼是估计器，不与LQR/LQI冲突。当前推荐“增广卡尔曼＋LQI＋IMU前馈”，ADRC作为对比方案；"
        "不建议同时运行一套高速ESO和卡尔曼扰动状态，以免对延迟视觉重复估计、产生振荡。",
        fill=LIGHT_BLUE,
    )

    doc.add_heading("6.2 1 cm精度与安全分级", level=2)
    add_table(
        doc,
        ["估计误差|e|", "车辆策略", "水管策略"],
        [
            ("<5 mm", "按地图正常限速运行", "精密LQI，日常建议不超过±4°"),
            ("5–8 mm", "减小加速度和加加速度", "提高回中优先级"),
            ("8–10 mm", "停止继续加速，准备受控停车", "允许短时使用±5.5°"),
            ("≥10 mm", "暂停地图任务", "回中/救球模式，硬限位±6°"),
            ("接近管端或反馈失联", "立即进入安全状态", "安全监督器覆盖普通LQI"),
        ],
        [1500, 3660, 4200],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body_paragraph(
        doc,
        "±6°能抵消的理想稳态轴向加速度约为g·tan6°≈1.03 m/s²，但还必须给位置反馈和车体俯仰留"
        "角度余量。因此正常地图规划建议先把车辆加减速度限制在0.4–0.6 m/s²，并限制加加速度；"
        "3 m/s²冲击超过机构稳态补偿能力，不是更换算法就能消除的。"
    )

    doc.add_heading("7 Simulink非线性仿真模型", level=1)
    add_body_paragraph(
        doc,
        "仿真采用“控制器—限幅/延迟—RS00与两连杆—钢球非线性对象—相机/IMU传感器—控制器”的闭环结构。"
        "固定步长为1/3000 s，使约30 Hz IMU源样本、200 Hz控制和100 Hz视觉落在公共仿真网格上。"
    )
    add_figure(
        doc,
        ROOT / "simulink_model_overview.png",
        "图3  自动生成的Simulink非线性闭环模型总览。正文按功能重新分组说明，原模型保留完整反馈与日志路径。",
        width=6.3,
    )
    add_table(
        doc,
        ["仿真模块", "主要内容", "对应文件"],
        [
            ("钢球对象", "平动/转动、Stribeck摩擦、滚阻、黏性/二次阻力、掉球判定", "sfun_ball_pipe_plant.m"),
            ("RS00与机构", "运动模式、力矩/速度/热限制、两连杆雅可比、变惯量和载荷", "sfun_rs00_joint_actuator.m"),
            ("相机", "100 Hz、量化、噪声、丢帧和延迟队列", "sfun_pipe_camera.m"),
            ("IMU", "约30 Hz唯一源样本、保持输出、偏置、噪声和处理延迟", "sfun_vehicle_imu.m"),
            ("控制器", "增广卡尔曼、LQI、车辆前馈、抗饱和和端部监督", "sfun_lqg_edge_controller.m"),
            ("车辆扰动", "轴向加速、俯仰、竖直振动和半正弦冲击", "vehicle_motion_profile.m"),
        ],
        [1600, 4840, 2920],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.8,
    )

    doc.add_heading("7.1 已纳入的现实因素", level=2)
    factors = [
        "钢球平动与转动惯量、滚动/滑动切换、静动摩擦过渡和滚动阻力；",
        "两连杆非线性角度关系、传动比、等效惯量、主动杆/蓝杆/水管质量和重力载荷；",
        "RS00位置刚度、速度阻尼、力矩/速度边界、电流环延迟、量化、回差和热降额；",
        "视觉帧率、噪声、像素量化、随机丢帧和树莓派处理/串口延迟；",
        "IMU采样、噪声、偏置、处理延迟以及车辆加减速、俯仰、振动和坑洼冲击；",
        "水管开放端掉球判定、机构有效性检查、角度/角速度限制和安全恢复。",
    ]
    for item in factors:
        add_list_item(doc, item, "•")

    doc.add_heading("7.2 尚未由实物确认的参数", level=2)
    add_callout(
        doc,
        "模型边界",
        "当前模型属于中等保真度工程模型，不是经过标定的数字孪生。球—管摩擦、水管质量与惯量、"
        "铰链间隙、支架刚度、真实RS00阶跃响应、视觉延迟分布和道路振动谱仍是主要不确定性。",
        fill=RISK,
        color="9B1C1C",
    )

    doc.add_heading("8 仿真结果与压力测试解释", level=1)
    doc.add_heading("8.1 标称跟踪仿真", level=2)
    add_figure(
        doc,
        ROOT / "ball_pipe_nominal_result.png",
        "图4  标称工况：参考位置、真实/估计位置与速度、水管角、滑移、RS00力矩和温度。"
        "该图用于验证闭环机理，不代表比赛1 cm稳态验收。",
        width=6.3,
    )
    add_body_paragraph(
        doc,
        "标称仿真中，卡尔曼估计能够跟随钢球真实位置和速度，水管命令受±6°限制，机构保持有效，"
        "RS00力矩远低于硬件峰值。参考位置包含±20 mm阶跃，用于验证新机构安全包线内的回中和跟踪能力。"
    )

    doc.add_heading("8.2 载车压力测试", level=2)
    add_figure(
        doc,
        ROOT / "vehicle_stress_test.png",
        "图5  低摩擦、车辆加减速、坑洼、视觉退化和超限3 m/s²脉冲的压力测试。"
        "红色虚线为管端，不是±10 mm精度线。",
        width=6.3,
    )
    stress_rows = [
        ("低摩擦静止", "7.381", "25.004", "0.671", "0.410", "是", "否"),
        ("正常车辆运动", "8.492", "25.000", "0.933", "0.651", "是", "否"),
        ("启动与急刹", "21.113", "48.118", "2.314", "0.651", "是", "否"),
        ("坑洼+视觉退化", "18.317", "40.571", "2.096", "0.727", "是", "否"),
        ("可行组合（优化）", "17.782", "45.729", "44.141", "0.701", "是", "否"),
        ("可行组合（仅KF-LQI）", "33.260", "88.038", "119.095", "0.693", "是", "否"),
        ("超新机构包线组合", "44.233", "92.709", "150.951", "0.713", "否", "否"),
        ("3.0 m/s²超限脉冲", "82.692", "115.015", "240.595", "0.694", "否", "否"),
    ]
    add_table(
        doc,
        ["场景", "RMS/mm", "峰值/mm", "滑移\nmm/s", "力矩\nN·m", "机械安全", "严格1 cm"],
        stress_rows,
        [2580, 1000, 1050, 1150, 1050, 1250, 1280],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 6,
        font_size=8.4,
    )
    add_callout(
        doc,
        "如何解读",
        "压力测试统一从+25 mm初始误差开始，因此若要求从t=0起始终小于10 mm，所有场景必然不合格。"
        "前六项“机械安全”为真，只表示未掉球、连杆有效、端部余量和滑移满足当前安全门槛。比赛报告应另行"
        "统计进入±10 mm所需时间，并在题目规定的评分区间检查峰值误差。",
        fill=CAUTION,
        color="7A5A00",
    )
    add_body_paragraph(
        doc,
        "在相同可行组合扰动下，加入车辆前馈和扰动估计后，RMS误差相对仅KF-LQI方案降低约46.5%，"
        "峰值滑移也由119.095降至44.141 mm/s。把WIT与MSPM0同时升级到115200 bit/s并获得"
        "200 Hz唯一IMU源样本的what-if结果为RMS 10.62 mm、峰值25.00 mm，较当前约30 Hz基线"
        "继续改善，但仍需硬件验证。3.0 m/s²脉冲产生115.015 mm峰值和240.595 mm/s滑移，说明"
        "该工况超出推荐物理包线。"
    )

    doc.add_heading("8.3 100、60与30 Hz视觉对照", level=2)
    add_figure(
        doc,
        ROOT / "camera_rate_comparison.png",
        "图6  相同200 Hz控制器和约30 Hz唯一IMU源样本条件下，100、60与30 Hz视觉的位置误差对照。"
        "测试从+5 mm开始，红线为±10 mm题目误差带。",
        width=6.3,
    )
    add_table(
        doc,
        ["工况", "视觉", "RMS/mm", "峰值/mm", "±10 mm占比", "严格通过"],
        [
            ("正常车辆运动", "100 Hz", "4.199", "8.701", "100%", "是"),
            ("正常车辆运动", "60 Hz", "4.853", "9.142", "100%", "是"),
            ("正常车辆运动", "30 Hz", "4.978", "7.781", "100%", "是"),
            ("启动、刹车与坑洼", "100 Hz", "48.173", "101.288", "33.72%", "否"),
            ("启动、刹车与坑洼", "60 Hz", "46.096", "97.453", "35.82%", "否"),
            ("启动、刹车与坑洼", "30 Hz", "49.321", "112.865", "34.54%", "否"),
        ],
        [2500, 1100, 1300, 1400, 1700, 1360],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 5,
        font_size=8.8,
    )
    add_callout(
        doc,
        "结论",
        "新机构和约30 Hz IMU源样本下，三种视觉帧率都能通过正常车辆场景，但启停/刹车/坑洼场景"
        "全部失败；帧率从30提高到100 Hz没有消除物理限角、机构限速和低速IMU前馈造成的包线限制。"
        "100 Hz仍是推荐视觉基线，但优化优先级应放在车辆加速度规划、115200/200 Hz IMU升级和"
        "RS00机构辨识，而不是继续盲目提高相机帧率。",
        fill=CAUTION,
        color="7A5A00",
    )
    add_body_paragraph(
        doc,
        "2026年7月31日模板匹配/投影标定版本部署后，钢球中部静置只读采样15.026 s，共1739帧；"
        "按capture_time_us去重仍为1739帧，found=true为1739/1739，有效率115.731 Hz。位置标准差"
        "0.25661 mm，相对中位数P95绝对偏差0.56805 mm，中心x标准差0.4677 px；处理耗时P95为"
        "1.442 ms。当前Simulink仍保守使用1.5 mm视觉噪声标准差，因为静态重复性不包含动态真值误差、"
        "曝光延迟和车辆振动。固定10 px半径、固定面积和临时confidence也不能作为正式质量门。"
    )

    doc.add_heading("9 英飞凌部署方案", level=1)
    add_body_paragraph(
        doc,
        "Simulink中的钢球、摩擦和车辆模型只用于仿真，不部署到单片机。部署到PSoC E84 M33的是"
        "状态估计、控制律、两连杆逆解、安全监督和RS00 CAN执行层。机械臂工程可只抽取已验证的RS00"
        "协议收发与反馈解析，不应移植其康复控制状态机。"
    )
    add_table(
        doc,
        ["任务", "频率/触发", "主要工作"],
        [
            ("RS00 CAN接收", "中断/DMA", "保存角度、速度、力矩、温度和接收时间戳"),
            ("MSPM0 IMU接收", "源事件约29 Hz", "按源序号接收姿态/轴向加速度和采样时间戳"),
            ("视觉接收", "帧事件≈100 Hz", "解析钢球位置、拍摄时间、帧号、置信度和CRC"),
            ("钢球控制", "200 Hz", "卡尔曼预测/延迟更新、LQI、前馈、抗饱和和限角"),
            ("RS00安全执行", "500 Hz", "目标插值、逆运动学、限速/力矩、反馈超时与CAN发送"),
            ("日志", "50–100 Hz", "记录原始传感器、估计状态、命令、反馈和故障标志"),
        ],
        [2000, 1800, 5560],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body_paragraph(
        doc,
        "建议把控制代码拆分为rs00_can、imu_link、vision_link、ball_kalman、ball_lqi、fourbar和"
        "ball_safety等无动态内存模块。三状态卡尔曼矩阵很小，可以使用float32手工实现，并用MATLAB"
        "生成逐帧输入输出测试向量；当前Level-2 MATLAB S-Function不适合直接代码生成。"
    )
    doc.add_heading("9.1 失联降级", level=2)
    degrade_items = [
        "相机失联<100 ms：继续状态预测；执行器饱和时冻结积分。",
        "相机失联100–250 ms：限制车辆加速度，水管指令缩小到±5°。",
        "相机失联>250 ms：车辆受控停车并进入保球模式。",
        "IMU唯一源样本超过100 ms：取消车辆前馈，禁止继续加速。",
        "RS00反馈超过10 ms：PSoC停止运动需求并进入电机安全状态。",
    ]
    for item in degrade_items:
        add_list_item(doc, item, "•")

    doc.add_heading("10 实物参数辨识与验证", level=1)
    add_body_paragraph(
        doc,
        "要让仿真能够预测物理世界，必须用实测数据识别参数，并使用未参与拟合的试验进行验证。"
        "建议按由低风险到高风险的顺序推进。"
    )
    validation_steps = [
        "机构标定：测量所有销轴中心距、C/O轴高度与水平距离；用电子水平仪记录−6°、−4°、0°、4°、6°时的RS00编码器角。",
        "摩擦辨识：缓慢增加倾角，记录开始运动角；在2°、4°、6°下记录钢球位置—时间曲线，每组至少重复10次。",
        "RS00辨识：无球条件下做±1°、±2°、±4°阶跃，记录命令、角度、速度、力矩和时间戳。",
        "视觉辨识：同时记录拍摄、处理完成和PSoC接收时间，统计位置误差、延迟分布和丢帧。",
        "硬件在环：PSoC运行真实200/500 Hz任务，由MATLAB注入可重复的视觉/IMU序列，逐帧比较结果。",
        "台架闭环：先限制水管±2°和约0.8 N·m，再逐步开放到±4°、±5.5°、±6°。",
        "车辆验证：从0.3 m/s²加速度开始，回放真实地图，最终以题目规定区间的峰值误差≤10 mm为准。",
    ]
    for index, item in enumerate(validation_steps, start=1):
        add_list_item(doc, item, f"{index}.")

    add_callout(
        doc,
        "验收原则",
        "算法正确、模型参数正确和硬件实现可靠是三个不同层次。只有仿真参数与实测吻合、硬件在环"
        "结果与MATLAB逐帧一致、并且多次实车测试均满足10 mm峰值要求，才能宣称系统达到题目指标。",
        fill=RISK,
        color="9B1C1C",
    )

    doc.add_heading("11 结论与当前设计基线", level=1)
    conclusions = [
        "本系统同时包含机械运动学、刚体动力学、传感器融合和控制理论，不能只归类为机械原理。",
        "当前推荐算法为延迟卡尔曼＋LQI＋IMU前馈＋两连杆逆解＋安全监督；ADRC用于对照，不与扰动卡尔曼重复叠加。",
        "当前实测机构为O轴38 mm、C轴93 mm、O-C水平285 mm、OA 35.0 mm、AB 55.5 mm、CB 300.1 mm，软件硬指令限角暂定±6°。",
        "PSoC E84直接控制RS00并承担最终安全；MSPM0发布带源序号的车身IMU，树莓派只提供视觉。",
        "当前压力测试可以发现机构和控制包线，但尚不能证明实物满足1 cm；摩擦、视觉延迟和RS00响应必须实测辨识。",
        "地图速度规划必须限制车辆加减速度和加加速度；正常起步建议先按0.4–0.6 m/s²以下设计，并由实车数据继续收紧。",
    ]
    for item in conclusions:
        add_list_item(doc, item, "•")

    doc.add_page_break()
    doc.add_heading("附录A 模型文件", level=1)
    files = [
        ("ball_pipe_nonlinear.slx", "自动生成的Simulink闭环模型"),
        ("ball_pipe_defaults.m", "全部机构、钢球、传感器、控制和仿真参数"),
        ("sfun_ball_pipe_plant.m", "钢球滚动/滑动非线性对象"),
        ("sfun_rs00_joint_actuator.m", "RS00、两连杆、负载、限幅和热模型"),
        ("sfun_lqg_edge_controller.m", "增广卡尔曼、LQI、车辆前馈和端部监督"),
        ("fourbar_kinematics.m", "两连杆正运动学、雅可比和有效性"),
        ("fourbar_inverse_kinematics.m", "水管角到RS00角的逆运动学"),
        ("run_ball_pipe_demo.m", "标称闭环仿真和曲线"),
        ("run_vehicle_stress_test.m", "载车低摩擦压力测试"),
        ("run_camera_rate_comparison.m", "100、60与30 Hz视觉的1 cm误差带对照"),
        ("analyze_two_link_mechanism.m", "工作区、传动比和等效惯量分析"),
        ("estimate_ball_pipe_parameters.m", "利用斜管实验估计钢球参数"),
    ]
    add_table(
        doc,
        ["文件", "作用"],
        files,
        [3500, 5860],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=9,
    )

    doc.add_heading("附录B 当前仿真关键参数", level=1)
    params = [
        ("钢球半径", "10 mm", "模型假设，待确认"),
        ("静/动摩擦系数", "标称0.050/0.030；压力测试最低0.020/0.012", "待实测"),
        ("水管最大角", "±6°", "新机构暂定硬指令限位"),
        ("水管角速度", "0.35 rad/s", "按最差雅可比和RS00限速保守设置"),
        ("控制/IMU唯一源样本/视觉", "200 / 约30 / 100 Hz", "当前模型"),
        ("视觉延迟", "25 ms处理 + 10 ms串口 = 35 ms", "模型假设"),
        ("IMU延迟", "35 ms", "当前WIT串行组+处理近似"),
        ("RS00指令/反馈延迟", "4.5 / 3.5 ms", "模型假设"),
        ("LQI现有增益", "[2.650, 0.910, 0.787]", "新限角下自动设计，仍须台架辨识"),
        ("机构安全工作区", "−6°到+6°", "最小|dθ/dq|=0.0467"),
    ]
    add_table(
        doc,
        ["参数", "当前值", "说明"],
        params,
        [2500, 3400, 3460],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.8,
    )
    add_callout(
        doc,
        "文档状态",
        "本报告记录当前设计与仿真基线。获得题目原文中的地图尺寸、速度/时间要求、评分时间窗，以及"
        "实测摩擦、视觉延迟和RS00响应数据后，应更新参数表、压力测试场景和1 cm合格判据。",
        fill=LIGHT_BLUE,
    )

    doc.core_properties.title = "钢球—水管两连杆平衡系统：机械原理、控制算法与Simulink非线性仿真"
    doc.core_properties.subject = "RS00两连杆钢球平衡系统技术说明"
    doc.core_properties.author = "项目组"
    doc.core_properties.keywords = "RS00, PSoC Edge E84, MSPM0G3507, Simulink, Kalman, LQI"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
